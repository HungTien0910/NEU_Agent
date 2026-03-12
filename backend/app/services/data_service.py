from __future__ import annotations

from io import BytesIO
from typing import Any, Callable, Dict, List, Optional

import pandas as pd
from openpyxl import load_workbook
from app.db.neo4j import get_neo4j_driver
from app.db.postgres import SessionLocal
from app.repositories.data_repo import (
    clear_all,
    list_academics,
    list_courses,
    list_lecturers,
    list_students,
    count_academics,
    count_courses,
    count_lecturers,
    count_students,
    list_academics_all,
    list_courses_all,
    list_lecturers_all,
    list_students_all,
    upsert_academics,
    upsert_courses,
    upsert_lecturers,
    upsert_students,
)
from app.models.academic_data import (
    AcademicRecord,
    CourseRecord,
    LecturerRecord,
    StudentRecord,
)
from app.services.import_progress import set_progress


PK_COLUMNS = {
    "Student": "MSSV",
    "Course": "MaHocPhan",
    "Lecturer": "MaGiangVien",
    "Academic": "RecordID",
}

REQUIRED_COLUMNS = {
    "Student": [
        "MSSV",
        "HoTen",
        "NgaySinh",
        "GioiTinh",
        "QueQuan",
        "Lop",
        "Khoa",
        "NamNhapHoc",
        "TrangThai",
        "GPA_Thang10",
        "GPA_Thang4",
    ],
    "Course": ["MaHocPhan", "TenMonHoc", "SoTinChi"],
    "Lecturer": ["MaGiangVien", "HoTen", "ChucDanh", "Khoa"],
    "Academic": [
        "RecordID",
        "MSSV",
        "MaHocPhan",
        "TenMonHoc",
        "SoTinChi",
        "HocKy",
        "NamHoc",
        "LopHocPhan",
        "GiangVienPhuTrach",
        "DiemThanhPhan",
        "DiemTongKet",
        "KetQua",
    ],
}


def read_excel(file_bytes: bytes) -> pd.ExcelFile:
    return pd.ExcelFile(BytesIO(file_bytes))


def validate_structure(file_bytes: bytes) -> Dict[str, Any]:
    excel = read_excel(file_bytes)
    missing_sheets = [s for s in REQUIRED_COLUMNS.keys() if s not in excel.sheet_names]
    missing_columns: List[Dict[str, Any]] = []

    for sheet, columns in REQUIRED_COLUMNS.items():
        if sheet not in excel.sheet_names:
            continue
        df = pd.read_excel(excel, sheet_name=sheet, nrows=1)
        sheet_cols = set(df.columns)
        missing = [col for col in columns if col not in sheet_cols]
        if missing:
            missing_columns.append({"sheet": sheet, "columns": missing})

    ok = not missing_sheets and not missing_columns
    detail_parts = []
    if missing_sheets:
        detail_parts.append(f"Thiếu sheet: {', '.join(missing_sheets)}")
    for item in missing_columns:
        cols = ", ".join(item["columns"])
        detail_parts.append(f"Sheet {item['sheet']} thiếu cột: {cols}")
    message = (
        "Cấu trúc file không chuẩn."
        if not detail_parts
        else "Cấu trúc file không chuẩn: " + " | ".join(detail_parts)
    )
    return {
        "ok": ok,
        "message": message,
        "missing_sheets": missing_sheets,
        "missing_columns": missing_columns,
        "sheets": excel.sheet_names,
    }


def build_preview(
    file_bytes: bytes,
    sheet: str,
    limit: int = 5,
    skip: int = 0,
    search: Optional[str] = None,
    file_name: Optional[str] = None,
) -> Dict[str, Any]:
    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    if not sheet_names:
        return {
            "file": {"name": file_name or "", "size": len(file_bytes)},
            "sheets": [],
            "sheet": sheet,
            "columns": [],
            "total_rows": 0,
            "rows": [],
            "limit": limit,
            "skip": skip,
            "page": 1,
            "total_pages": 1,
            "summary": [],
        }
    if sheet not in sheet_names:
        sheet = sheet_names[0]
    ws = wb[sheet]

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    headers = [str(h) if h is not None else "" for h in (header_row or [])]

    total_rows = max(ws.max_row - 1, 0)
    if skip < 0:
        skip = 0
    rows: List[Dict[str, Any]] = []
    search_value = (search or "").strip().lower()
    pk_column = PK_COLUMNS.get(sheet)
    pk_index = headers.index(pk_column) if pk_column in headers else None

    if search_value and pk_index is not None:
        matched = 0
        collected = 0
        for row in ws.iter_rows(
            min_row=2,
            max_row=ws.max_row,
            max_col=len(headers),
            values_only=True,
        ):
            cell = row[pk_index] if pk_index < len(row) else None
            cell_text = "" if cell is None else str(cell).strip().lower()
            if search_value in cell_text:
                if matched >= skip and collected < limit:
                    item: Dict[str, Any] = {}
                    for idx, key in enumerate(headers):
                        value = row[idx] if idx < len(row) else ""
                        item[key] = "" if value is None else value
                    rows.append(item)
                    collected += 1
                matched += 1
            if collected >= limit and matched >= skip + limit:
                # continue counting for total rows
                continue
        total_rows = matched
    else:
        if skip > total_rows:
            skip = total_rows
        start_row = 2 + skip
        end_row = min(start_row + max(limit, 0) - 1, ws.max_row)
        if headers and start_row <= end_row:
            for row in ws.iter_rows(
                min_row=start_row,
                max_row=end_row,
                max_col=len(headers),
                values_only=True,
            ):
                item: Dict[str, Any] = {}
                for idx, key in enumerate(headers):
                    value = row[idx] if idx < len(row) else ""
                    item[key] = "" if value is None else value
                rows.append(item)

    summary: List[Dict[str, Any]] = []
    for name in sheet_names:
        s_ws = wb[name]
        s_header = next(
            s_ws.iter_rows(min_row=1, max_row=1, values_only=True), None
        )
        s_cols = [str(h) if h is not None else "" for h in (s_header or [])]
        summary.append(
            {
                "sheet": name,
                "rows": max(s_ws.max_row - 1, 0),
                "columns": s_cols,
                "columns_count": len(s_cols),
            }
        )
    wb.close()
    return {
        "file": {"name": file_name or "", "size": len(file_bytes)},
        "sheets": sheet_names,
        "sheet": sheet,
        "columns": headers,
        "total_rows": int(total_rows),
        "rows": rows,
        "limit": limit,
        "skip": skip,
        "page": int(skip / limit) + 1 if limit else 1,
        "total_pages": int((total_rows + limit - 1) / limit) if limit else 1,
        "summary": summary,
    }


def validate_excel(file_bytes: bytes, limit_errors: int = 200) -> Dict[str, Any]:
    excel = read_excel(file_bytes)
    errors: List[Dict[str, Any]] = []
    total = 0
    valid = 0

    for sheet in excel.sheet_names:
        df = pd.read_excel(excel, sheet_name=sheet)
        total += df.shape[0]
        pk_col = PK_COLUMNS.get(sheet)
        if not pk_col or pk_col not in df.columns:
            continue
        null_mask = df[pk_col].isna() | (df[pk_col].astype(str).str.strip() == "")
        invalid_rows = df[null_mask]
        valid += df.shape[0] - invalid_rows.shape[0]

        for idx, row in invalid_rows.iterrows():
            if len(errors) >= limit_errors:
                break
            errors.append(
                {
                    "sheet": sheet,
                    "row": int(idx) + 2,
                    "column": pk_col,
                    "type": "NULL",
                    "value": "",
                }
            )

    return {
        "total": total,
        "valid": valid,
        "errors": len(errors),
        "details": errors,
        "note": "Thiếu khóa chính (PK) ở một số dòng",
    }


def import_to_neo4j(
    file_bytes: bytes,
    progress_cb: Optional[Callable[[int, str], None]] = None,
) -> None:
    excel = read_excel(file_bytes)
    driver = get_neo4j_driver()

    with driver.session() as session:
        if progress_cb:
            progress_cb(15, "Neo4j: Student")
        _import_students(session, excel)
        if progress_cb:
            progress_cb(35, "Neo4j: Course")
        _import_courses(session, excel)
        if progress_cb:
            progress_cb(50, "Neo4j: Lecturer")
        _import_lecturers(session, excel)
        if progress_cb:
            progress_cb(65, "Neo4j: Academic")
        _import_academic(session, excel)
    driver.close()


def save_to_postgres(
    db,
    file_bytes: bytes,
    progress_cb: Optional[Callable[[int, str], None]] = None,
) -> None:
    excel = read_excel(file_bytes)

    if "Student" in excel.sheet_names:
        df = pd.read_excel(excel, sheet_name="Student").fillna("")
        rows = []
        for _, row in df.iterrows():
            mssv = _safe_value(row.get("MSSV"))
            if not mssv:
                continue
            rows.append(
                {
                    "mssv": mssv,
                    "ho_ten": _safe_value(row.get("HoTen")),
                    "ngay_sinh": _safe_value(row.get("NgaySinh")),
                    "gioi_tinh": _safe_value(row.get("GioiTinh")),
                    "que_quan": _safe_value(row.get("QueQuan")),
                    "lop": _safe_value(row.get("Lop")),
                    "khoa": _safe_value(row.get("Khoa")),
                    "nam_nhap_hoc": _safe_value(row.get("NamNhapHoc")),
                    "trang_thai": _safe_value(row.get("TrangThai")),
                    "gpa10": _safe_float(row.get("GPA_Thang10")),
                    "gpa4": _safe_float(row.get("GPA_Thang4")),
                }
            )
        upsert_students(db, rows)
        if progress_cb:
            progress_cb(75, "PostgreSQL: Student")

    if "Course" in excel.sheet_names:
        df = pd.read_excel(excel, sheet_name="Course").fillna("")
        rows = []
        for _, row in df.iterrows():
            code = _safe_value(row.get("MaHocPhan"))
            if not code:
                continue
            rows.append(
                {
                    "ma_hoc_phan": code,
                    "ten_mon_hoc": _safe_value(row.get("TenMonHoc")),
                    "so_tin_chi": _safe_float(row.get("SoTinChi")),
                }
            )
        upsert_courses(db, rows)
        if progress_cb:
            progress_cb(82, "PostgreSQL: Course")

    if "Lecturer" in excel.sheet_names:
        df = pd.read_excel(excel, sheet_name="Lecturer").fillna("")
        rows = []
        for _, row in df.iterrows():
            code = _safe_value(row.get("MaGiangVien"))
            if not code:
                continue
            rows.append(
                {
                    "ma_giang_vien": code,
                    "ho_ten": _safe_value(row.get("HoTen")),
                    "chuc_danh": _safe_value(row.get("ChucDanh")),
                    "khoa": _safe_value(row.get("Khoa")),
                }
            )
        upsert_lecturers(db, rows)
        if progress_cb:
            progress_cb(88, "PostgreSQL: Lecturer")

    if "Academic" in excel.sheet_names:
        df = pd.read_excel(excel, sheet_name="Academic").fillna("")
        rows = []
        for _, row in df.iterrows():
            record_id = _safe_value(row.get("RecordID"))
            if not record_id:
                continue
            rows.append(
                {
                    "record_id": record_id,
                    "mssv": _safe_value(row.get("MSSV")),
                    "ma_hoc_phan": _safe_value(row.get("MaHocPhan")),
                    "ten_mon_hoc": _safe_value(row.get("TenMonHoc")),
                    "so_tin_chi": _safe_float(row.get("SoTinChi")),
                    "hoc_ky": _safe_value(row.get("HocKy")),
                    "nam_hoc": _safe_value(row.get("NamHoc")),
                    "lop_hoc_phan": _safe_value(row.get("LopHocPhan")),
                    "giang_vien_phu_trach": _safe_value(row.get("GiangVienPhuTrach")),
                    "diem_thanh_phan": _safe_float(row.get("DiemThanhPhan")),
                    "diem_tong_ket": _safe_float(row.get("DiemTongKet")),
                    "ket_qua": _safe_value(row.get("KetQua")),
                }
            )
        upsert_academics(db, rows)
        if progress_cb:
            progress_cb(95, "PostgreSQL: Academic")


def run_import_job(job_id: str, file_bytes: bytes) -> None:
    try:
        set_progress(job_id, "running", 5, "Đang đẩy dữ liệu")
        structure = validate_structure(file_bytes)
        if not structure["ok"]:
            set_progress(job_id, "error", 100, "Lỗi cấu trúc file", structure["message"])
            return

        def _cb(percent: int, step: str) -> None:
            set_progress(job_id, "running", percent, "Đang đẩy dữ liệu")

        set_progress(job_id, "running", 10, "Đang đẩy dữ liệu")
        import_to_neo4j(file_bytes, progress_cb=_cb)
        set_progress(job_id, "running", 70, "Đang đẩy dữ liệu")
        db = SessionLocal()
        try:
            save_to_postgres(db, file_bytes, progress_cb=_cb)
        finally:
            db.close()
        set_progress(job_id, "done", 100, "Hoàn tất import")
    except Exception as exc:  # pragma: no cover
        set_progress(job_id, "error", 100, "Lỗi import", str(exc))


def clear_all_data() -> None:
    driver = get_neo4j_driver()
    with driver.session() as session:
        session.run("MATCH (n) DETACH DELETE n")
    driver.close()


def _import_students(session, excel: pd.ExcelFile) -> None:
    if "Student" not in excel.sheet_names:
        return
    df = pd.read_excel(excel, sheet_name="Student")
    df = df.fillna("")
    df["MSSV"] = df["MSSV"].apply(_normalize_code)
    df = df[df["MSSV"] != ""]
    rows = [
        {key: _normalize_cell(value) for key, value in row.items()}
        for row in df.to_dict(orient="records")
    ]

    query = """
    UNWIND $rows AS row
    MERGE (s:Student {mssv: row.MSSV})
    SET s.hoTen = row.HoTen,
        s.ngaySinh = row.NgaySinh,
        s.gioiTinh = row.GioiTinh,
        s.queQuan = row.QueQuan,
        s.lop = row.Lop,
        s.khoa = row.Khoa,
        s.namNhapHoc = row.NamNhapHoc,
        s.trangThai = row.TrangThai,
        s.gpa10 = row.GPA_Thang10,
        s.gpa4 = row.GPA_Thang4
    MERGE (d:Department {tenKhoa: row.Khoa})
    MERGE (c:Class {tenLop: row.Lop})
    MERGE (s)-[:BELONGS_TO]->(d)
    MERGE (s)-[:IN_CLASS]->(c)
    """
    session.run(query, rows=rows)


def _import_courses(session, excel: pd.ExcelFile) -> None:
    if "Course" not in excel.sheet_names:
        return
    df = pd.read_excel(excel, sheet_name="Course")
    df = df.fillna("")
    df["MaHocPhan"] = df["MaHocPhan"].apply(_normalize_code)
    df = df[df["MaHocPhan"] != ""]
    rows = [
        {key: _normalize_cell(value) for key, value in row.items()}
        for row in df.to_dict(orient="records")
    ]
    query = """
    UNWIND $rows AS row
    MERGE (c:Course {maHocPhan: row.MaHocPhan})
    SET c.tenMonHoc = row.TenMonHoc,
        c.soTinChi = row.SoTinChi
    """
    session.run(query, rows=rows)


def _import_lecturers(session, excel: pd.ExcelFile) -> None:
    if "Lecturer" not in excel.sheet_names:
        return
    df = pd.read_excel(excel, sheet_name="Lecturer")
    df = df.fillna("")
    df["MaGiangVien"] = df["MaGiangVien"].apply(_normalize_code)
    df = df[df["MaGiangVien"] != ""]
    rows = [
        {key: _normalize_cell(value) for key, value in row.items()}
        for row in df.to_dict(orient="records")
    ]
    query = """
    UNWIND $rows AS row
    MERGE (l:Lecturer {maGiangVien: row.MaGiangVien})
    SET l.hoTen = row.HoTen,
        l.chucDanh = row.ChucDanh,
        l.khoa = row.Khoa
    MERGE (d:Department {tenKhoa: row.Khoa})
    MERGE (l)-[:IN_DEPARTMENT]->(d)
    """
    session.run(query, rows=rows)


def _import_academic(session, excel: pd.ExcelFile) -> None:
    if "Academic" not in excel.sheet_names:
        return
    df = pd.read_excel(excel, sheet_name="Academic")
    df = df.fillna("")
    df["RecordID"] = df["RecordID"].apply(_normalize_code)
    df["MSSV"] = df["MSSV"].apply(_normalize_code)
    df["MaHocPhan"] = df["MaHocPhan"].apply(_normalize_code)
    df = df[df["RecordID"] != ""]
    rows = [
        {key: _normalize_cell(value) for key, value in row.items()}
        for row in df.to_dict(orient="records")
    ]
    query = """
    UNWIND $rows AS row
    OPTIONAL MATCH (s:Student {mssv: row.MSSV})
    OPTIONAL MATCH (l:Lecturer {hoTen: row.GiangVienPhuTrach})
    MERGE (r:AcademicRecord {recordId: row.RecordID})
    SET r.mssv = row.MSSV,
        r.maHocPhan = row.MaHocPhan,
        r.tenMonHoc = row.TenMonHoc,
        r.soTinChi = row.SoTinChi,
        r.hocKy = row.HocKy,
        r.namHoc = row.NamHoc,
        r.lopHocPhan = row.LopHocPhan,
        r.diemThanhPhan = row.DiemThanhPhan,
        r.diemTongKet = row.DiemTongKet,
        r.ketQua = row.KetQua,
        r.giangVienPhuTrach = row.GiangVienPhuTrach
    FOREACH (_ IN CASE WHEN row.MaHocPhan IS NULL OR row.MaHocPhan = "" THEN [] ELSE [1] END |
        MERGE (c:Course {maHocPhan: row.MaHocPhan})
        SET c.tenMonHoc = row.TenMonHoc,
            c.soTinChi = row.SoTinChi
        MERGE (r)-[:OF_COURSE]->(c)
    )
    FOREACH (_ IN CASE WHEN s IS NULL THEN [] ELSE [1] END |
        MERGE (s)-[:HAS_RECORD]->(r)
    )
    FOREACH (_ IN CASE WHEN l IS NULL THEN [] ELSE [1] END |
        MERGE (r)-[:TAUGHT_BY]->(l)
    )
    """
    session.run(query, rows=rows)


def fetch_sheet_items(
    db, sheet: str, search: Optional[str] = None, limit: int = 50, skip: int = 0
) -> List[Dict[str, Any]]:
    if sheet == "Student":
        items = list_students(db, search, limit, skip)
        return [
            {
                "id": item.id,
                "mssv": item.mssv,
                "hoTen": item.ho_ten,
                "ngaySinh": item.ngay_sinh,
                "gioiTinh": item.gioi_tinh,
                "queQuan": item.que_quan,
                "lop": item.lop,
                "khoa": item.khoa,
                "namNhapHoc": item.nam_nhap_hoc,
                "trangThai": item.trang_thai,
                "gpa10": item.gpa10,
                "gpa4": item.gpa4,
            }
            for item in items
        ]
    if sheet == "Course":
        items = list_courses(db, search, limit, skip)
        return [
            {
                "id": item.id,
                "maHocPhan": item.ma_hoc_phan,
                "tenMonHoc": item.ten_mon_hoc,
                "soTinChi": item.so_tin_chi,
            }
            for item in items
        ]
    if sheet == "Lecturer":
        items = list_lecturers(db, search, limit, skip)
        return [
            {
                "id": item.id,
                "maGiangVien": item.ma_giang_vien,
                "hoTen": item.ho_ten,
                "chucDanh": item.chuc_danh,
                "khoa": item.khoa,
            }
            for item in items
        ]
    items = list_academics(db, search, limit, skip)
    return [
        {
            "id": item.id,
            "recordId": item.record_id,
            "mssv": item.mssv,
            "maHocPhan": item.ma_hoc_phan,
            "tenMonHoc": item.ten_mon_hoc,
            "soTinChi": item.so_tin_chi,
            "hocKy": item.hoc_ky,
            "namHoc": item.nam_hoc,
            "lopHocPhan": item.lop_hoc_phan,
            "giangVienPhuTrach": item.giang_vien_phu_trach,
            "diemThanhPhan": item.diem_thanh_phan,
            "diemTongKet": item.diem_tong_ket,
            "ketQua": item.ket_qua,
        }
        for item in items
    ]


def count_sheet_items(db, sheet: str, search: Optional[str] = None) -> int:
    if sheet == "Student":
        return count_students(db, search)
    if sheet == "Course":
        return count_courses(db, search)
    if sheet == "Lecturer":
        return count_lecturers(db, search)
    return count_academics(db, search)


def _serialize_student(item: StudentRecord) -> Dict[str, Any]:
    return {
        "id": item.id,
        "mssv": item.mssv,
        "hoTen": item.ho_ten,
        "ngaySinh": item.ngay_sinh,
        "gioiTinh": item.gioi_tinh,
        "queQuan": item.que_quan,
        "lop": item.lop,
        "khoa": item.khoa,
        "namNhapHoc": item.nam_nhap_hoc,
        "trangThai": item.trang_thai,
        "gpa10": item.gpa10,
        "gpa4": item.gpa4,
    }


def _serialize_course(item: CourseRecord) -> Dict[str, Any]:
    return {
        "id": item.id,
        "maHocPhan": item.ma_hoc_phan,
        "tenMonHoc": item.ten_mon_hoc,
        "soTinChi": item.so_tin_chi,
    }


def _serialize_lecturer(item: LecturerRecord) -> Dict[str, Any]:
    return {
        "id": item.id,
        "maGiangVien": item.ma_giang_vien,
        "hoTen": item.ho_ten,
        "chucDanh": item.chuc_danh,
        "khoa": item.khoa,
    }


def _serialize_academic(item: AcademicRecord) -> Dict[str, Any]:
    return {
        "id": item.id,
        "recordId": item.record_id,
        "mssv": item.mssv,
        "maHocPhan": item.ma_hoc_phan,
        "tenMonHoc": item.ten_mon_hoc,
        "soTinChi": item.so_tin_chi,
        "hocKy": item.hoc_ky,
        "namHoc": item.nam_hoc,
        "lopHocPhan": item.lop_hoc_phan,
        "giangVienPhuTrach": item.giang_vien_phu_trach,
        "diemThanhPhan": item.diem_thanh_phan,
        "diemTongKet": item.diem_tong_ket,
        "ketQua": item.ket_qua,
    }


def update_sheet_item(db, sheet: str, item_id: int, payload: Dict[str, Any]) -> Dict[str, Any]:
    model_map = {
        "Student": StudentRecord,
        "Course": CourseRecord,
        "Lecturer": LecturerRecord,
        "Academic": AcademicRecord,
    }
    field_map = {
        "Student": {
            "hoTen": "ho_ten",
            "ngaySinh": "ngay_sinh",
            "gioiTinh": "gioi_tinh",
            "queQuan": "que_quan",
            "lop": "lop",
            "khoa": "khoa",
            "namNhapHoc": "nam_nhap_hoc",
            "trangThai": "trang_thai",
            "gpa10": "gpa10",
            "gpa4": "gpa4",
        },
        "Course": {"tenMonHoc": "ten_mon_hoc", "soTinChi": "so_tin_chi"},
        "Lecturer": {"hoTen": "ho_ten", "chucDanh": "chuc_danh", "khoa": "khoa"},
        "Academic": {
            "mssv": "mssv",
            "maHocPhan": "ma_hoc_phan",
            "tenMonHoc": "ten_mon_hoc",
            "soTinChi": "so_tin_chi",
            "hocKy": "hoc_ky",
            "namHoc": "nam_hoc",
            "lopHocPhan": "lop_hoc_phan",
            "giangVienPhuTrach": "giang_vien_phu_trach",
            "diemThanhPhan": "diem_thanh_phan",
            "diemTongKet": "diem_tong_ket",
            "ketQua": "ket_qua",
        },
    }
    serializer_map = {
        "Student": _serialize_student,
        "Course": _serialize_course,
        "Lecturer": _serialize_lecturer,
        "Academic": _serialize_academic,
    }
    model = model_map.get(sheet)
    if not model:
        raise ValueError("Sheet không hợp lệ.")
    record = db.query(model).filter(model.id == item_id).first()
    if not record:
        raise ValueError("Không tìm thấy dữ liệu.")
    float_fields = {
        "Student": {"gpa10", "gpa4"},
        "Course": {"soTinChi"},
        "Academic": {"soTinChi", "diemThanhPhan", "diemTongKet"},
    }
    mapping = field_map.get(sheet, {})
    for key, attr in mapping.items():
        if key in payload:
            value = payload[key]
            if key in float_fields.get(sheet, set()):
                if value in (None, ""):
                    value = None
                else:
                    try:
                        value = float(value)
                    except (TypeError, ValueError):
                        raise ValueError(f"Giá trị {key} không hợp lệ.")
            setattr(record, attr, value)
    db.commit()
    db.refresh(record)
    _sync_update_neo4j(sheet, record)
    return serializer_map[sheet](record)


def create_sheet_item(db, sheet: str, payload: Dict[str, Any]) -> Dict[str, Any]:
    model_map = {
        "Student": StudentRecord,
        "Course": CourseRecord,
        "Lecturer": LecturerRecord,
        "Academic": AcademicRecord,
    }
    field_map = {
        "Student": {
            "mssv": "mssv",
            "hoTen": "ho_ten",
            "ngaySinh": "ngay_sinh",
            "gioiTinh": "gioi_tinh",
            "queQuan": "que_quan",
            "lop": "lop",
            "khoa": "khoa",
            "namNhapHoc": "nam_nhap_hoc",
            "trangThai": "trang_thai",
            "gpa10": "gpa10",
            "gpa4": "gpa4",
        },
        "Course": {
            "maHocPhan": "ma_hoc_phan",
            "tenMonHoc": "ten_mon_hoc",
            "soTinChi": "so_tin_chi",
        },
        "Lecturer": {
            "maGiangVien": "ma_giang_vien",
            "hoTen": "ho_ten",
            "chucDanh": "chuc_danh",
            "khoa": "khoa",
        },
        "Academic": {
            "recordId": "record_id",
            "mssv": "mssv",
            "maHocPhan": "ma_hoc_phan",
            "tenMonHoc": "ten_mon_hoc",
            "soTinChi": "so_tin_chi",
            "hocKy": "hoc_ky",
            "namHoc": "nam_hoc",
            "lopHocPhan": "lop_hoc_phan",
            "giangVienPhuTrach": "giang_vien_phu_trach",
            "diemThanhPhan": "diem_thanh_phan",
            "diemTongKet": "diem_tong_ket",
            "ketQua": "ket_qua",
        },
    }
    pk_map = {
        "Student": "mssv",
        "Course": "ma_hoc_phan",
        "Lecturer": "ma_giang_vien",
        "Academic": "record_id",
    }
    serializer_map = {
        "Student": _serialize_student,
        "Course": _serialize_course,
        "Lecturer": _serialize_lecturer,
        "Academic": _serialize_academic,
    }
    model = model_map.get(sheet)
    if not model:
        raise ValueError("Sheet không hợp lệ.")
    mapping = field_map.get(sheet, {})
    pk_field = pk_map.get(sheet)
    if not pk_field:
        raise ValueError("Không tìm thấy khóa chính.")

    float_fields = {
        "Student": {"gpa10", "gpa4"},
        "Course": {"soTinChi"},
        "Academic": {"soTinChi", "diemThanhPhan", "diemTongKet"},
    }

    data: Dict[str, Any] = {}
    for key, attr in mapping.items():
        if key in payload:
            value = payload[key]
            if key in float_fields.get(sheet, set()):
                if value in (None, ""):
                    value = None
                else:
                    try:
                        value = float(value)
                    except (TypeError, ValueError):
                        raise ValueError(f"Giá trị {key} không hợp lệ.")
            data[attr] = value

    if pk_field not in data or not data[pk_field]:
        raise ValueError("Thiếu khóa chính.")
    exists = db.query(model).filter(getattr(model, pk_field) == data[pk_field]).first()
    if exists:
        raise ValueError("Khóa chính đã tồn tại.")

    record = model(**data)
    db.add(record)
    db.commit()
    db.refresh(record)
    _sync_update_neo4j(sheet, record)
    return serializer_map[sheet](record)


def delete_sheet_item(db, sheet: str, item_id: int) -> None:
    model_map = {
        "Student": StudentRecord,
        "Course": CourseRecord,
        "Lecturer": LecturerRecord,
        "Academic": AcademicRecord,
    }
    model = model_map.get(sheet)
    if not model:
        raise ValueError("Sheet không hợp lệ.")
    record = db.query(model).filter(model.id == item_id).first()
    if not record:
        raise ValueError("Không tìm thấy dữ liệu.")
    _sync_delete_neo4j(sheet, record)
    db.delete(record)
    db.commit()


def export_sheet_excel(db, sheet: str, search: Optional[str] = None) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet

    if sheet == "Student":
        headers = REQUIRED_COLUMNS["Student"]
        records = list_students_all(db, search)
        ws.append(headers)
        for item in records:
            ws.append(
                [
                    item.mssv,
                    item.ho_ten,
                    item.ngay_sinh,
                    item.gioi_tinh,
                    item.que_quan,
                    item.lop,
                    item.khoa,
                    item.nam_nhap_hoc,
                    item.trang_thai,
                    item.gpa10,
                    item.gpa4,
                ]
            )
    elif sheet == "Course":
        headers = REQUIRED_COLUMNS["Course"]
        records = list_courses_all(db, search)
        ws.append(headers)
        for item in records:
            ws.append([item.ma_hoc_phan, item.ten_mon_hoc, item.so_tin_chi])
    elif sheet == "Lecturer":
        headers = REQUIRED_COLUMNS["Lecturer"]
        records = list_lecturers_all(db, search)
        ws.append(headers)
        for item in records:
            ws.append([item.ma_giang_vien, item.ho_ten, item.chuc_danh, item.khoa])
    else:
        headers = REQUIRED_COLUMNS["Academic"]
        records = list_academics_all(db, search)
        ws.append(headers)
        for item in records:
            ws.append(
                [
                    item.record_id,
                    item.mssv,
                    item.ma_hoc_phan,
                    item.ten_mon_hoc,
                    item.so_tin_chi,
                    item.hoc_ky,
                    item.nam_hoc,
                    item.lop_hoc_phan,
                    item.giang_vien_phu_trach,
                    item.diem_thanh_phan,
                    item.diem_tong_ket,
                    item.ket_qua,
                ]
            )

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def export_all_excel(db) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()

    # Student
    ws = wb.active
    ws.title = "Student"
    headers = REQUIRED_COLUMNS["Student"]
    ws.append(headers)
    for item in list_students_all(db, None):
        ws.append(
            [
                item.mssv,
                item.ho_ten,
                item.ngay_sinh,
                item.gioi_tinh,
                item.que_quan,
                item.lop,
                item.khoa,
                item.nam_nhap_hoc,
                item.trang_thai,
                item.gpa10,
                item.gpa4,
            ]
        )

    # Course
    ws = wb.create_sheet("Course")
    headers = REQUIRED_COLUMNS["Course"]
    ws.append(headers)
    for item in list_courses_all(db, None):
        ws.append([item.ma_hoc_phan, item.ten_mon_hoc, item.so_tin_chi])

    # Lecturer
    ws = wb.create_sheet("Lecturer")
    headers = REQUIRED_COLUMNS["Lecturer"]
    ws.append(headers)
    for item in list_lecturers_all(db, None):
        ws.append([item.ma_giang_vien, item.ho_ten, item.chuc_danh, item.khoa])

    # Academic
    ws = wb.create_sheet("Academic")
    headers = REQUIRED_COLUMNS["Academic"]
    ws.append(headers)
    for item in list_academics_all(db, None):
        ws.append(
            [
                item.record_id,
                item.mssv,
                item.ma_hoc_phan,
                item.ten_mon_hoc,
                item.so_tin_chi,
                item.hoc_ky,
                item.nam_hoc,
                item.lop_hoc_phan,
                item.giang_vien_phu_trach,
                item.diem_thanh_phan,
                item.diem_tong_ket,
                item.ket_qua,
            ]
        )

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _sync_update_neo4j(sheet: str, record: Any) -> None:
    driver = get_neo4j_driver()
    with driver.session() as session:
        if sheet == "Student":
            session.run(
                """
                MERGE (s:Student {mssv: $mssv})
                SET s.hoTen = $hoTen,
                    s.ngaySinh = $ngaySinh,
                    s.gioiTinh = $gioiTinh,
                    s.queQuan = $queQuan,
                    s.khoa = $khoa,
                    s.namNhapHoc = $namNhapHoc,
                    s.trangThai = $trangThai,
                    s.gpa10 = $gpa10,
                    s.gpa4 = $gpa4
                MERGE (d:Department {tenKhoa: $khoa})
                MERGE (c:Class {tenLop: $lop})
                MERGE (s)-[:BELONGS_TO]->(d)
                MERGE (s)-[:IN_CLASS]->(c)
                """,
                {
                    "mssv": record.mssv,
                    "hoTen": record.ho_ten,
                    "ngaySinh": record.ngay_sinh,
                    "gioiTinh": record.gioi_tinh,
                    "queQuan": record.que_quan,
                    "namNhapHoc": record.nam_nhap_hoc,
                    "trangThai": record.trang_thai,
                    "gpa10": record.gpa10,
                    "gpa4": record.gpa4,
                    "khoa": record.khoa,
                    "lop": record.lop,
                },
            )
        elif sheet == "Course":
            session.run(
                """
                MERGE (c:Course {maHocPhan: $maHocPhan})
                SET c.tenMonHoc = $tenMonHoc,
                    c.soTinChi = $soTinChi
                """,
                {
                    "maHocPhan": record.ma_hoc_phan,
                    "tenMonHoc": record.ten_mon_hoc,
                    "soTinChi": record.so_tin_chi,
                },
            )
        elif sheet == "Lecturer":
            session.run(
                """
                MERGE (l:Lecturer {maGiangVien: $maGiangVien})
                SET l.hoTen = $hoTen,
                    l.chucDanh = $chucDanh,
                    l.khoa = $khoa
                MERGE (d:Department {tenKhoa: $khoa})
                MERGE (l)-[:IN_DEPARTMENT]->(d)
                """,
                {
                    "maGiangVien": record.ma_giang_vien,
                    "hoTen": record.ho_ten,
                    "chucDanh": record.chuc_danh,
                    "khoa": record.khoa,
                },
            )
        else:
            session.run(
                """
                MERGE (r:AcademicRecord {recordId: $recordId})
                SET r.maHocPhan = $maHocPhan,
                    r.hocKy = $hocKy,
                    r.namHoc = $namHoc,
                    r.lopHocPhan = $lopHocPhan,
                    r.diemThanhPhan = $diemThanhPhan,
                    r.diemTongKet = $diemTongKet,
                    r.ketQua = $ketQua
                WITH r
                MATCH (s:Student {mssv: $mssv})
                MERGE (s)-[:HAS_RECORD]->(r)
                MERGE (c:Course {maHocPhan: $maHocPhan})
                MERGE (r)-[:OF_COURSE]->(c)
                MERGE (l:Lecturer {hoTen: $giangVienPhuTrach})
                MERGE (r)-[:TAUGHT_BY]->(l)
                """,
                {
                    "recordId": record.record_id,
                    "maHocPhan": record.ma_hoc_phan,
                    "mssv": record.mssv,
                    "hocKy": record.hoc_ky,
                    "namHoc": record.nam_hoc,
                    "lopHocPhan": record.lop_hoc_phan,
                    "diemThanhPhan": record.diem_thanh_phan,
                    "diemTongKet": record.diem_tong_ket,
                    "ketQua": record.ket_qua,
                    "giangVienPhuTrach": record.giang_vien_phu_trach,
                },
            )
    driver.close()


def _sync_delete_neo4j(sheet: str, record: Any) -> None:
    driver = get_neo4j_driver()
    with driver.session() as session:
        if sheet == "Student":
            session.run(
                "MATCH (s:Student {mssv: $mssv}) DETACH DELETE s",
                {"mssv": record.mssv},
            )
        elif sheet == "Course":
            session.run(
                "MATCH (c:Course {maHocPhan: $maHocPhan}) DETACH DELETE c",
                {"maHocPhan": record.ma_hoc_phan},
            )
        elif sheet == "Lecturer":
            session.run(
                "MATCH (l:Lecturer {maGiangVien: $maGiangVien}) DETACH DELETE l",
                {"maGiangVien": record.ma_giang_vien},
            )
        else:
            session.run(
                "MATCH (r:AcademicRecord {recordId: $recordId}) DETACH DELETE r",
                {"recordId": record.record_id},
            )
    driver.close()


def clear_postgres(db) -> None:
    clear_all(db)


def _safe_value(value: Any) -> Optional[str]:
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.date().isoformat()
    return str(value).strip() if value is not None else None


def _normalize_code(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    if isinstance(value, (int,)):
        return str(value).strip()
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value)).strip()
        return str(value).strip()
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text


def _normalize_cell(value: Any) -> Any:
    if value is None or pd.isna(value):
        return ""
    if isinstance(value, pd.Timestamp):
        return value.date().isoformat()
    if hasattr(value, "item"):
        try:
            return value.item()
        except Exception:
            return value
    return value


def _safe_float(value: Any) -> Optional[float]:
    if pd.isna(value) or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None
